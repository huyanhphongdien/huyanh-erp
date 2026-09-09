# -*- coding: utf-8 -*-
"""Sinh file Word HƯỚNG DẪN CÂN MỦ LẺ (cho nhân viên cân).
   Chạy: python docs/generate_huong_dan_can_mu_le.py
   Ra:   docs/HUONG_DAN_CAN_MU_LE.docx
"""
import sys, os
sys.stdout.reconfigure(encoding="utf-8")
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

GREEN = RGBColor(0x1B, 0x4D, 0x3E)
GOLD  = RGBColor(0xB4, 0x53, 0x09)
RED   = RGBColor(0xB9, 0x1C, 0x1C)
MUTED = RGBColor(0x5B, 0x65, 0x60)

doc = Document()
for s in doc.sections:
    s.top_margin = s.bottom_margin = Inches(0.6)
    s.left_margin = s.right_margin = Inches(0.7)

normal = doc.styles["Normal"]
normal.font.name = "Segoe UI"
normal.font.size = Pt(11.5)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def run(p, text, bold=False, color=None, size=None):
    r = p.add_run(text)
    r.bold = bold
    if color: r.font.color.rgb = color
    if size: r.font.size = Pt(size)
    return r


def heading(n, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    run(p, f"{n} · ", bold=True, color=GREEN, size=14)
    run(p, title, bold=True, color=GREEN, size=14)
    return p


def steps(items):
    for it in items:
        p = doc.add_paragraph(style="List Number")
        # cho phép ("bình thường", [("Bấm ", False),("Kết nối", True)]) hoặc chuỗi thường
        if isinstance(it, list):
            for txt, bold in it:
                run(p, txt, bold=bold)
        else:
            run(p, it)


def callout(title, body, kind="warn"):
    """Hộp cảnh báo: 1 ô bảng có nền màu."""
    fill = "FFF7E6" if kind == "warn" else "FDECEA"
    tcolor = GOLD if kind == "warn" else RED
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    cell = t.cell(0, 0)
    shade(cell, fill)
    p = cell.paragraphs[0]
    run(p, title + "  ", bold=True, color=tcolor)
    # body: chuỗi hoặc list (txt,bold)
    if isinstance(body, list):
        for txt, bold in body:
            run(p, txt, bold=bold)
    else:
        run(p, body)
    doc.add_paragraph()


# ─── ĐẦU TRANG ────────────────────────────────────────────────────────────────
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
run(title, "HƯỚNG DẪN CÂN MỦ LẺ", bold=True, color=GREEN, size=20)
sub = doc.add_paragraph()
run(sub, "Cao su Huy Anh — dành cho nhân viên cân   ·   Mở app tại: ", color=MUTED, size=10.5)
run(sub, "canle.huyanhrubber.vn", bold=True, color=GREEN, size=10.5)

lead = doc.add_paragraph()
run(lead, "Quy trình mua mủ tạp của hộ dân / khách vãng lai. Cân từng bao → chờ cán bộ đo DRC → "
          "nhập DRC & đơn giá → in phiếu. ")
run(lead, "Tiền do phòng kế toán chi qua Đề nghị thanh toán, KHÔNG đưa tiền tại cân.", bold=True)

flow = doc.add_paragraph()
run(flow, "Tóm tắt:  ", bold=True, color=GREEN)
run(flow, "1 Đăng nhập  →  2 Nối cân  →  3 Cân từng bao  →  4 Lưu chờ DRC  →  5 Nhập DRC + giá → In",
    color=MUTED)

# ─── 1 ────────────────────────────────────────────────────────────────────────
heading(1, "Mở app & đăng nhập")
steps([
    [("Mở trình duyệt ", False), ("Chrome", True), (" (hoặc Edge) trên máy tính trạm cân.", False)],
    [("Vào địa chỉ ", False), ("canle.huyanhrubber.vn", True), (" (nên lưu Bookmark cho lần sau).", False)],
    [("Nhập ", False), ("mã PIN", True), (" của bạn để đăng nhập.", False)],
])
callout("⚠ Phải dùng Chrome hoặc Edge:",
        "Trình duyệt khác (Firefox, Safari…) không đọc được đầu cân.")

# ─── 2 ────────────────────────────────────────────────────────────────────────
heading(2, "Kết nối đầu cân (làm 1 lần đầu ngày)")
steps([
    [("Góc phải trên nếu ghi ", False), ("○ Chưa nối", True), (" thì cần kết nối.", False)],
    [("Bấm ", False), ("Kết nối cổng COM", True), (" → chọn đúng cổng của đầu cân → Kết nối.", False)],
    [("Thành công: góc phải hiện ", False), ("● Cân OK", True),
     (" và số cân nhảy theo vật đặt lên.", False)],
])
callout("Đầu cân tự dò thông số:",
        [("Không cần chỉnh gì thủ công; nếu lỡ sai app tự dò lại. Máy nhớ cổng nên ", False),
         ("các lần sau thường tự nối", True), (", khỏi làm lại bước này.", False)])

# ─── 3 ────────────────────────────────────────────────────────────────────────
heading(3, "Cân khách mới — cân từng bao")
steps([
    [("Trang chủ bấm ", False), ("＋ Cân khách mới", True), (".", False)],
    [("Nhập ", False), ("Tên khách", True), (" (bắt buộc); có SĐT thì nhập thêm.", False)],
    [("Đặt 1 bao lên cân → chờ chữ ", False), ("● Đã đứng số", True),
     (" (đèn xanh) → bấm ", False), ("⚡ LẤY SỐ → THÊM BAO", True), (".", False)],
    [("Nhấc bao xuống, đặt bao tiếp theo, lặp lại cho hết các bao.", False)],
    [("Kiểm tra Tổng khối lượng đúng với số bao đã cân.", False)],
    [("Bấm ", False), ("💾 LƯU — CHỜ DRC", True), (".", False)],
])
callout("💡 Mua CẢ BÌ — không trừ bì:",
        [("Số cân trên màn = ", False), ("số thực nhận", True),
         (" (đã tính cả bao), không phải trừ gì thêm.", False)])
callout("⚠ Chờ đèn xanh mới bấm lấy số:",
        "Số còn nhảy mà bấm là sai. Nhấc bao trước xuống rồi mới cân bao sau.")

# ─── 4 ────────────────────────────────────────────────────────────────────────
heading(4, "Lấy mẫu → chờ đo DRC")
steps([
    "Sau khi lưu, phiếu vào khối ⏳ Chờ DRC ở Trang chủ.",
    [("Lấy ", False), ("mẫu từng bao, gộp chung", True), (" → đưa cán bộ đo DRC (1 kết quả cho cả xe).", False)],
    "Trong lúc chờ, cân tiếp khách khác bình thường — nhiều xe chờ DRC cùng lúc được.",
])

# ─── 5 ────────────────────────────────────────────────────────────────────────
heading(5, "Chốt DRC + In phiếu")
steps([
    [("Có kết quả DRC: Trang chủ → khối ⏳ Chờ DRC → bấm đúng phiếu khách đó (", False),
     ("Nhập DRC →", True), (").", False)],
    [("Nhập ", False), ("DRC (%)", True), (" cán bộ báo.", False)],
    [("Nhập ", False), ("Đơn giá", True), (" (₫/kg khô) theo giá ngày.", False)],
    [("App tự tính Khô = Tổng kg × DRC%, rồi Thành tiền = Khô × Đơn giá. Xem lại số tiền.", False)],
    [("Bấm ", False), ("🖨 CHỐT & IN PHIẾU", True), (" → phiếu ra máy in, đưa khách.", False)],
])
callout("Ví dụ tính tiền:",
        [("Tổng 55 kg · DRC 43% → khô 23,65 kg · giá 30.000 ₫/kg khô → ", False),
         ("Thành tiền = 709.500 ₫.", True)])

# ─── 6 ────────────────────────────────────────────────────────────────────────
heading(6, "Lưu ý & xử lý sự cố")
rows = [
    ("Tình huống", "Cách xử lý"),
    ('Góc phải ghi "Chưa nối"', "Bấm Kết nối cổng COM → chọn cổng (bước 2)."),
    ("Số cân không lên / đứng im", "Kiểm tra dây đầu cân ↔ máy tính. Vẫn không được → báo kỹ thuật."),
    ("Số cân lệch với đầu cân", "Báo kỹ thuật ngay, đừng tự cân tiếp (sai tiền khách)."),
    ("Bấm Lưu bị lỗi mạng", 'Nếu app báo "đã tạo phiếu, đừng lưu lại" thì KHÔNG bấm lưu lần 2 — báo kế toán.'),
    ("Cân nhầm / khách bỏ về", "Vào phiếu → Huỷ phiếu (ghi lý do). Phiếu đã chốt/in thì báo kế toán."),
]
tbl = doc.add_table(rows=len(rows), cols=2)
tbl.style = "Table Grid"
tbl.columns[0].width = Inches(2.3); tbl.columns[1].width = Inches(4.4)
for i, (a, b) in enumerate(rows):
    c0, c1 = tbl.cell(i, 0), tbl.cell(i, 1)
    if i == 0:
        shade(c0, "EFEEE8"); shade(c1, "EFEEE8")
        run(c0.paragraphs[0], a, bold=True); run(c1.paragraphs[0], b, bold=True)
    else:
        run(c0.paragraphs[0], a, bold=True, color=GREEN); run(c1.paragraphs[0], b)

doc.add_paragraph()
callout("⛔ TUYỆT ĐỐI:",
        [("Không đưa tiền cho khách tại cân — tiền do kế toán chi qua Đề nghị thanh toán. "
          "Không bấm Lưu / Chốt hai lần cho cùng một xe.", True)], kind="stop")

foot = doc.add_paragraph()
foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
run(foot, "Cao su Huy Anh — Hướng dẫn Cân mủ lẻ · Vướng mắc kỹ thuật: báo IT / kế toán.",
    color=MUTED, size=9.5)

out = os.path.join(os.path.dirname(__file__), "HUONG_DAN_CAN_MU_LE.docx")
doc.save(out)
print("Đã tạo:", out)
